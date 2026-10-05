from datetime import datetime
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from core.models import Departamento, Equipamiento, HistorialEquipo, Usuario
from core.services.asignacion_equipos import source_digest


class AssignEquipmentUsersCommandTests(TestCase):
    def setUp(self):
        self.department = Departamento.objects.create(nombre='Asignación Equipos')
        self.active_user = Usuario.objects.create(
            nombre_completo='Usuario Activo',
            usuario_red='usuario.activo',
            correo_corp='activo@example.com',
            hostname='cl_activo',
            estado='ACTIVO',
            departamento=self.department,
        )
        self.inactive_user = Usuario.objects.create(
            nombre_completo='Usuario Baja',
            usuario_red='usuario.baja',
            correo_corp='baja@example.com',
            hostname='cl_baja',
            estado='BAJA',
            departamento=self.department,
        )
        self.assignable = Equipamiento.objects.create(
            tipo='Notebook',
            marca='HP',
            modelo='840',
            numero_serie='SERIE-ACTIVA',
            estado='STOCK',
        )
        self.blocked = Equipamiento.objects.create(
            tipo='Celular',
            marca='Samsung',
            modelo='A1',
            numero_serie='SERIE-BAJA',
            estado='STOCK',
        )

    def test_assigns_only_safe_match_and_writes_pending_report(self):
        with TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory) / 'Usuarios Simi.xlsx'
            report = Path(temporary_directory) / 'pendientes.txt'
            book = Workbook()
            sheet = book.active
            sheet.cell(2, 2, 'NOMBRES Y APELLIDOS')
            sheet.cell(4, 2, 'Usuario Activo')
            sheet.cell(4, 3, 'HP')
            sheet.cell(4, 4, '840')
            sheet.cell(4, 5, 'SERIE-ACTIVA')
            sheet.cell(4, 7, datetime(2026, 9, 1))
            sheet.cell(5, 2, 'Usuario Baja')
            sheet.cell(5, 8, 'Samsung')
            sheet.cell(5, 9, 'A1')
            sheet.cell(5, 10, 'SERIE-BAJA')
            book.save(source)

            call_command(
                'asignar_equipos_usuarios',
                '--source', str(source),
                '--apply',
                '--expected-sha256', source_digest(source),
                '--pending-report', str(report),
                stdout=StringIO(),
            )

            self.assignable.refresh_from_db()
            self.blocked.refresh_from_db()
            self.assertEqual(self.assignable.usuario, self.active_user)
            self.assertEqual(self.assignable.estado, 'ASIGNADO')
            self.assertEqual(self.assignable.hostname, 'cl_activo')
            self.assertEqual(str(self.assignable.fecha_asignacion), '2026-09-01')
            self.assertIsNone(self.blocked.usuario)
            self.assertEqual(self.blocked.estado, 'STOCK')
            self.assertIn('USUARIO_EN_ESTADO_BAJA', report.read_text(encoding='utf-8-sig'))
            self.assertTrue(
                HistorialEquipo.objects.filter(
                    equipo=self.assignable,
                    modificado_por='importacion_equipos_v11',
                ).exists()
            )
