"""Planificación determinista de la carga v1.1; no escribe ni registra secretos."""

import hashlib
import re
from collections import Counter, defaultdict
from datetime import datetime
from ipaddress import ip_address

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook

from core.models import (
    Anexo, Departamento, Equipamiento, IP, PCGenerico, PerfilGenerico, Usuario,
    _normalize_key,
)
from core.serializers import IP_ALLOWED_NETWORKS


FILES = {
    'usuarios': 'Usuarios Simi.xlsx',
    'ips': 'BARRIDO IP OFICIAL.xlsx',
    'anexos': 'Anexos.xlsx',
    'perfiles': 'Perfiles Genericos.xlsx',
    'pcs': 'PCs Genericos.xlsx',
}
EMPTY = {'', '-', 'N/A', 'N/I', 'FALTA', 'S/C', 'NAN', 'NONE'}
HOSTNAME = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._-]*$')


def cell(row, column):
    raw = row[column - 1] if len(row) >= column else None
    text = ' '.join(str(raw).strip().split()) if raw is not None else ''
    return '' if text.upper() in EMPTY else text


def normalized(text):
    return ' '.join(text.casefold().split())


def valid_email(text):
    if not text or len(text) > 254:
        return False
    try:
        validate_email(text)
    except DjangoValidationError:
        return False
    return True


def secret_fits(text):
    # Fernet + prefijo ENC2:: deben caber en CharField(max_length=255).
    return len(text.encode('utf-8')) <= 127


def data_rows(sheet, header_row):
    for number, row in enumerate(sheet.values, 1):
        if number > header_row:
            yield number, row


def check_header(sheet, row_number, column, expected):
    header = next(sheet.iter_rows(
        min_row=row_number, max_row=row_number, values_only=True
    ))
    if expected not in cell(header, column).upper():
        raise ValueError(
            f'Encabezado no reconocido en {sheet.title}, fila {row_number}, '
            f'columna {column}; esperado: {expected}'
        )


def file_digest(root):
    digest = hashlib.sha256()
    for name in FILES.values():
        digest.update(name.encode('utf-8'))
        with (root / name).open('rb') as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b''):
                digest.update(chunk)
    return digest.hexdigest()


class ImportPlan:
    """Solo inserta identidades únicas; todo conflicto queda como fila pendiente."""

    def __init__(self):
        self.counts = Counter()
        self.issues = []
        self.departments = {}
        self.users = []
        self.equipment = []
        self.unassigned_equipment = []
        self.ips = []
        self.annexes = []
        self.profiles = []
        self.pcs = []
        self.user_index = {name: defaultdict(set) for name in ('name', 'email')}

    def issue(self, module, sheet, row, code):
        self.counts[code] += 1
        self.issues.append({
            'module': module, 'sheet': sheet, 'row': row, 'code': code,
        })

    def public_report(self, digest):
        return {
            'input_sha256': digest,
            'planned': {
                'departamentos': len([d for d in self.departments.values() if d['new']]),
                'usuarios': len(self.users),
                'equipos': len(self.equipment),
                'ips': len(self.ips),
                'anexos': len(self.annexes),
                'perfiles': len(self.profiles),
                'pcs': len(self.pcs),
            },
            'counts': dict(sorted(self.counts.items())),
            'issues': self.issues,
            'policy': {
                'subareas': 'vacías; Departamento/Área se carga como Departamento',
                'ips': 'solo dirección, estado LIBRE, sin observación ni asignación',
                'anexo_usuario': 'solo coincidencia exacta; otros sin usuario',
                'equipos': 'todos se crean sin usuario para enroque manual',
                'gmail': 'se añade @gmail.com a las cuentas que no incluyen dominio',
                'perfil_tipo_ausente': 'On Premise',
                'existing_records': 'se conservan sin cambios',
                'secret_values_in_report': False,
            },
            'equipos_sin_usuario': self.unassigned_equipment,
        }

    def department(self, name, module, sheet, row):
        if not name or len(name) > 100:
            self.issue(module, sheet, row, 'DEPARTAMENTO_INVALIDO')
            return None
        identity = _normalize_key(name)
        if identity in self.departments:
            return identity
        existing = Departamento.objects.filter(nombre_normalizado=identity).first()
        if existing and not existing.activo:
            self.issue(module, sheet, row, 'DEPARTAMENTO_INACTIVO')
            return None
        self.departments[identity] = {'name': name, 'new': existing is None, 'id': existing.pk if existing else None}
        return identity

    def read_users(self, sheet):
        check_header(sheet, 3, 5, 'NOMBRE DE USUARIO')
        source = []
        for number, row in data_rows(sheet, 3):
            if not any(cell(row, c) for c in (2, 3, 5, 7, 13, 20, 29, 37)):
                continue
            source.append((number, row))
        usernames = Counter(normalized(cell(row, 5)) for _, row in source if cell(row, 5))
        emails = Counter(normalized(cell(row, 7)) for _, row in source if cell(row, 7))
        names = Counter(normalized(cell(row, 2)) for _, row in source if cell(row, 2))
        existing_users = {normalized(u.usuario_red): u for u in Usuario.objects.all()}
        existing_emails = {normalized(u.correo_corp): u for u in Usuario.objects.all()}
        existing_names = {normalized(u.nombre_completo) for u in Usuario.objects.all()}
        for number, row in source:
            name, dept, username, email = (cell(row, c) for c in (2, 3, 5, 7))
            uk, ek, nk = normalized(username), normalized(email), normalized(name)
            if not (name and username and email and dept):
                self.issue('usuarios', sheet.title, number, 'USUARIO_INCOMPLETO')
                continue
            if (len(name) > 150 or len(username) > 50 or len(email) > 254
                    or any(c.isspace() for c in username) or not valid_email(email)):
                self.issue('usuarios', sheet.title, number, 'USUARIO_FORMATO_INVALIDO')
                continue
            if usernames[uk] > 1 or emails[ek] > 1 or names[nk] > 1:
                self.issue('usuarios', sheet.title, number, 'USUARIO_DUPLICADO_ARCHIVO')
                continue
            by_username, by_email = existing_users.get(uk), existing_emails.get(ek)
            if by_username or by_email:
                if by_username and by_email and by_username.pk == by_email.pk:
                    reference = ('db', by_username.pk)
                    self.counts['usuarios_existentes_sin_cambios'] += 1
                else:
                    self.issue('usuarios', sheet.title, number, 'USUARIO_CONFLICTO_BD')
                    continue
            elif nk in existing_names:
                self.issue('usuarios', sheet.title, number, 'NOMBRE_EXISTE_EN_BD')
                continue
            else:
                dept_key = self.department(dept, 'usuarios', sheet.title, number)
                if dept_key is None:
                    continue
                hostname = cell(row, 10)
                if hostname and (len(hostname) > 50 or not HOSTNAME.fullmatch(hostname)):
                    self.issue('usuarios', sheet.title, number, 'HOSTNAME_INVALIDO_OMITIDO')
                    hostname = ''
                gmail = cell(row, 8)
                gmail_password = cell(row, 9)
                if gmail and '@' not in gmail:
                    gmail += '@gmail.com'
                    self.counts['gmail_dominio_completado'] += 1
                if gmail and not valid_email(gmail):
                    self.issue('usuarios', sheet.title, number, 'GMAIL_NO_ES_CORREO_OMITIDO')
                    gmail = ''
                    gmail_password = ''
                if not gmail and gmail_password:
                    self.issue('usuarios', sheet.title, number, 'GMAIL_PASSWORD_SIN_CUENTA_OMITIDO')
                    gmail_password = ''
                if gmail_password and not secret_fits(gmail_password):
                    self.issue('usuarios', sheet.title, number, 'GMAIL_PASSWORD_LARGO_OMITIDO')
                    gmail_password = ''
                reference = ('new', number)
                self.users.append({
                    'row': number, 'department': dept_key,
                    'data': {
                        'nombre_completo': name, 'usuario_red': username,
                        'correo_corp': email, 'cargo': cell(row, 4) or None,
                        'hostname': hostname or None, 'gmail': gmail or None,
                        'password_gmail': gmail_password or None,
                        'sif': cell(row, 11).casefold() in {'si', 'sí'},
                        'vpn_cisco': cell(row, 12).casefold() in {'si', 'sí'},
                    },
                })
            self.user_index['name'][nk].add(reference)
            self.user_index['email'][ek].add(reference)
        self.read_equipment(sheet, source)

    def read_equipment(self, sheet, source):
        specs = (
            ('Notebook', 13, 14, 15, 16, 17, None, None),
            ('Celular', 20, 21, 22, 24, 27, 23, 26),
            ('Tablet', 29, 30, 31, 32, 35, None, 34),
            ('BAM / Router', 37, 38, 39, 40, 42, None, None),
        )
        candidates = []
        for number, row in source:
            for kind, brand_c, model_c, serial_c, af_c, date_c, phone_c, pin_c in specs:
                brand, model, serial = (cell(row, c) for c in (brand_c, model_c, serial_c))
                if not any((brand, model, serial)):
                    continue
                candidates.append((number, row, kind, brand, model, serial, af_c, date_c, phone_c, pin_c))
        serials = Counter(normalized(item[5]) for item in candidates if item[5])
        afs = Counter(normalized(cell(item[1], item[6])) for item in candidates if cell(item[1], item[6]))
        existing_serials = {normalized(x) for x in Equipamiento.objects.exclude(numero_serie__isnull=True).values_list('numero_serie', flat=True)}
        existing_afs = {normalized(x) for x in Equipamiento.objects.exclude(af__isnull=True).values_list('af', flat=True)}
        for number, row, kind, brand, model, serial, af_c, date_c, phone_c, pin_c in candidates:
            af = cell(row, af_c)
            if not brand or not model or not (serial or af) or len(brand) > 50 or len(model) > 50 or len(serial) > 20:
                self.issue('equipos', sheet.title, number, 'EQUIPO_IDENTIFICADOR_INVALIDO')
                continue
            if serial and (serials[normalized(serial)] > 1 or normalized(serial) in existing_serials):
                self.issue('equipos', sheet.title, number, 'EQUIPO_SERIE_DUPLICADA')
                continue
            if af and (len(af) > 12 or not af.isalnum() or afs[normalized(af)] > 1 or normalized(af) in existing_afs):
                if not serial:
                    self.issue('equipos', sheet.title, number, 'EQUIPO_ACTIVO_FIJO_CONFLICTO')
                    continue
                self.issue('equipos', sheet.title, number, 'EQUIPO_ACTIVO_FIJO_CONFLICTO_OMITIDO')
                af = ''
            phone = cell(row, phone_c) if phone_c else ''
            if phone and len(phone) == 11 and phone.startswith('56') and phone.isdigit():
                phone = '+' + phone
                self.counts['equipos_telefono_normalizado'] += 1
            if phone and not (len(phone) == 12 and phone.startswith('+') and phone[1:].isdigit()):
                self.issue('equipos', sheet.title, number, 'EQUIPO_TELEFONO_INVALIDO_OMITIDO')
                phone = ''
            pin = cell(row, pin_c) if pin_c else ''
            raw_pin = row[pin_c - 1] if pin_c and len(row) >= pin_c else None
            if pin and not isinstance(raw_pin, str):
                self.issue('equipos', sheet.title, number, 'EQUIPO_PIN_NUMERICO_REVISAR')
            if pin and not secret_fits(pin):
                self.issue('equipos', sheet.title, number, 'EQUIPO_PIN_LARGO_OMITIDO')
                pin = ''
            raw_date = row[date_c - 1] if len(row) >= date_c else None
            date = raw_date.date() if isinstance(raw_date, datetime) else None
            self.unassigned_equipment.append({
                'sheet': sheet.title, 'row': number, 'tipo': kind,
                'marca': brand, 'modelo': model, 'serie': serial or None,
                'activo_fijo': af or None,
            })
            self.equipment.append({
                'row': number, 'user': None,
                'data': {
                    'tipo': kind, 'marca': brand, 'modelo': model,
                    'numero_serie': serial or None, 'af': af or None,
                    'numero_telefono': phone or None, 'pin': pin or None,
                    'fecha_asignacion': date,
                },
            })

    def read_ips(self, book):
        existing = set(IP.objects.values_list('direccion_ip', flat=True))
        seen = set()
        for sheet in book:
            header = 4 if sheet.title.startswith('172.') else 2
            check_header(sheet, header, 2, 'DIRECCION IP')
            for number, row in enumerate(sheet.values, 1):
                raw = cell(row, 2)
                if not raw:
                    continue
                try:
                    parsed = ip_address(raw)
                except ValueError:
                    continue
                network = next((n for n in IP_ALLOWED_NETWORKS if parsed in n), None)
                if not network or parsed in (network.network_address, network.broadcast_address):
                    self.issue('ips', sheet.title, number, 'IP_NO_ASIGNABLE')
                    continue
                ip = str(parsed)
                if ip in seen or ip in existing:
                    self.issue('ips', sheet.title, number, 'IP_DUPLICADA_O_EXISTENTE')
                    continue
                seen.add(ip)
                self.ips.append({'row': number, 'sheet': sheet.title, 'ip': ip})

    def read_annexes(self, sheet):
        check_header(sheet, 2, 5, 'ANEXO')
        source = [(n, r) for n, r in data_rows(sheet, 2) if cell(r, 5)]
        numbers = Counter(normalized(cell(row, 5)) for _, row in source)
        existing = {normalized(x) for x in Anexo.objects.values_list('numero_anexo', flat=True)}
        assigned = set(Anexo.objects.exclude(usuario__isnull=True).values_list('usuario_id', flat=True))
        planned_assignments = set()
        for number, row in source:
            extension = cell(row, 5)
            if not extension.isdigit() or len(extension) > 10 or numbers[normalized(extension)] > 1 or normalized(extension) in existing:
                self.issue('anexos', sheet.title, number, 'ANEXO_NUMERO_INVALIDO_O_EXISTENTE')
                continue
            exterior, observation = cell(row, 6), cell(row, 8)
            if len(exterior) > 30:
                self.issue('anexos', sheet.title, number, 'ANEXO_EXTERIOR_LARGO_OMITIDO')
                exterior = ''
            name, email = cell(row, 2), cell(row, 7)
            matches = []
            if name:
                matches.append(self.user_index['name'].get(normalized(name), set()))
            if email:
                matches.append(self.user_index['email'].get(normalized(email), set()))
            common = set.intersection(*matches) if matches else set()
            user = next(iter(common)) if len(common) == 1 else None
            if user and (user in planned_assignments or (user[0] == 'db' and user[1] in assigned)):
                user = None
                self.issue('anexos', sheet.title, number, 'ANEXO_USUARIO_YA_TIENE_ANEXO')
            if user:
                planned_assignments.add(user)
            elif matches:
                self.issue('anexos', sheet.title, number, 'ANEXO_SIN_USUARIO_EXACTO')
            self.annexes.append({
                'row': number, 'user': user, 'data': {
                    'numero_anexo': extension, 'exterior': exterior or None,
                    'observaciones': observation or None,
                },
            })

    def read_profiles(self, sheet):
        check_header(sheet, 3, 3, 'NOMBRE DE USUARIO')
        source = [(n, r) for n, r in data_rows(sheet, 3) if any(cell(r, c) for c in (2, 3, 5, 6))]
        names = Counter(normalized(cell(row, 3)) for _, row in source if cell(row, 3))
        existing = {normalized(x) for x in PerfilGenerico.objects.values_list('usuario', flat=True)}
        for number, row in source:
            name, username, email, dept, note = (cell(row, c) for c in (2, 3, 5, 6, 7))
            if not name or not username or not dept or len(name) > 150 or len(username) > 100 or any(c.isspace() for c in username):
                self.issue('perfiles', sheet.title, number, 'PERFIL_INCOMPLETO_O_INVALIDO')
                continue
            if names[normalized(username)] > 1 or normalized(username) in existing:
                self.issue('perfiles', sheet.title, number, 'PERFIL_DUPLICADO_O_EXISTENTE')
                continue
            if email and not valid_email(email):
                self.issue('perfiles', sheet.title, number, 'PERFIL_CORREO_INVALIDO_OMITIDO')
                email = ''
            if 'office 365' in note.casefold() or 'o365' in note.casefold():
                kind = 'O365'
            elif 'on-premise' in note.casefold() or 'on premise' in note.casefold():
                kind = 'On Premise'
            else:
                kind = 'On Premise'
                self.counts['perfiles_tipo_predeterminado'] += 1
            dept_key = self.department(dept, 'perfiles', sheet.title, number)
            if dept_key is None:
                continue
            password = cell(row, 4)
            if password and not secret_fits(password):
                self.issue('perfiles', sheet.title, number, 'PERFIL_PASSWORD_LARGO_OMITIDO')
                password = ''
            self.profiles.append({
                'row': number, 'department': dept_key,
                'data': {
                    'nombre': name, 'usuario': username, 'correo': email or None,
                    'password': password or None, 'tipo': kind,
                    'observaciones': note or None,
                },
            })

    def read_pcs(self, sheet):
        check_header(sheet, 4, 5, 'HOSTNAME')
        source = [(n, r) for n, r in data_rows(sheet, 4) if cell(r, 3) or cell(r, 5)]
        hostnames = Counter(normalized(cell(row, 5)) for _, row in source if cell(row, 5))
        serials = Counter(normalized(cell(row, 9)) for _, row in source if cell(row, 9))
        afs = Counter(normalized(cell(row, 10)) for _, row in source if cell(row, 10))
        existing_hosts = {normalized(x) for x in PCGenerico.objects.values_list('hostname', flat=True)}
        existing_serials = {normalized(x) for x in PCGenerico.objects.exclude(numero_serie__isnull=True).values_list('numero_serie', flat=True)}
        existing_afs = {normalized(x) for x in PCGenerico.objects.exclude(activo_fijo__isnull=True).values_list('activo_fijo', flat=True)}
        for number, row in source:
            username, hostname, serial = (cell(row, c) for c in (3, 5, 9))
            department = cell(row, 6)
            if not username or len(username) > 150 or not hostname or len(hostname) > 100 or not HOSTNAME.fullmatch(hostname):
                self.issue('pcs', sheet.title, number, 'PC_IDENTIFICADOR_INVALIDO')
                continue
            if not department:
                self.issue('pcs', sheet.title, number, 'PC_DEPARTAMENTO_REQUERIDO')
                continue
            if hostnames[normalized(hostname)] > 1 or normalized(hostname) in existing_hosts:
                self.issue('pcs', sheet.title, number, 'PC_HOSTNAME_DUPLICADO_O_EXISTENTE')
                continue
            if serial and (len(serial) > 20 or serials[normalized(serial)] > 1 or normalized(serial) in existing_serials):
                self.issue('pcs', sheet.title, number, 'PC_SERIE_INVALIDA_O_DUPLICADA')
                continue
            af = cell(row, 10)
            if af and (len(af) > 12 or not af.isalnum() or afs[normalized(af)] > 1 or normalized(af) in existing_afs):
                self.issue('pcs', sheet.title, number, 'PC_ACTIVO_FIJO_INVALIDO_OMITIDO')
                af = ''
            password = cell(row, 4)
            if password and not secret_fits(password):
                self.issue('pcs', sheet.title, number, 'PC_PASSWORD_LARGO_OMITIDO')
                password = ''
            department_key = self.department(
                department,
                'pcs',
                sheet.title,
                number,
            )
            self.pcs.append({
                'row': number,
                'department': department_key,
                'data': {
                    'usuario_local': username, 'hostname': hostname,
                    'password': password or None,
                    'dpto_area': department,
                    'marca': cell(row, 7) or None, 'modelo': cell(row, 8) or None,
                    'numero_serie': serial or None, 'activo_fijo': af or None,
                    'observaciones': cell(row, 11) or None,
                },
            })


def make_plan(root):
    plan = ImportPlan()
    books = {}
    try:
        for module, filename in FILES.items():
            books[module] = load_workbook(root / filename, read_only=True, data_only=True)
        plan.read_users(books['usuarios'].active)
        plan.read_ips(books['ips'])
        plan.read_annexes(books['anexos'].active)
        plan.read_profiles(books['perfiles'].active)
        plan.read_pcs(books['pcs'].active)
    finally:
        for book in books.values():
            book.close()
    return plan
