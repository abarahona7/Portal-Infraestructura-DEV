import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from cryptography.fernet import Fernet
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from openpyxl import Workbook

from core.crypto import decrypt_val
from core.models import Anexo, Departamento, Equipamiento, IP, PCGenerico, PerfilGenerico, Usuario
from core.services.importacion_v11 import file_digest


@override_settings(FIELD_ENCRYPTION_KEY=Fernet.generate_key().decode())
class ImportV11Tests(TestCase):
    def test_legacy_importer_is_permanently_retired(self):
        with self.assertRaisesRegex(CommandError, 'fue retirado'):
            call_command('importar_datos', stdout=io.StringIO())

    def make_workbooks(self, root):
        specs = [
            ('Usuarios Simi.xlsx', 3, {5: 'NOMBRE DE USUARIO'}, [
                (4, {2: 'Persona Prueba', 3: 'TI', 5: 'prueba',
                     7: 'prueba@example.com', 8: 'aliasprueba',
                     9: 'CLAVE_MUY_SECRETA', 13: 'Marca A',
                     14: 'Modelo A', 15: 'SERIE_OK'}),
                (5, {2: 'Persona Sin Red', 3: 'TI', 13: 'Marca',
                     14: 'Modelo', 15: 'SERIE123'}),
            ]),
            ('BARRIDO IP OFICIAL.xlsx', 2, {2: 'DIRECCION IP'}, [
                (3, {2: '192.168.10.10', 3: 'En uso', 4: 'USUARIO'}),
            ]),
            ('Anexos.xlsx', 2, {5: 'ANEXO'}, [
                (3, {2: 'Persona Prueba', 5: '4000', 7: 'prueba@example.com'}),
                (4, {2: 'Sin coincidencia', 5: '4001'}),
            ]),
            ('Perfiles Genericos.xlsx', 3, {3: 'NOMBRE DE USUARIO'}, [
                (4, {2: 'Prueba', 3: 'perfil_prueba', 4: 'OTRA_CLAVE_SECRETA',
                     6: 'TI'}),
            ]),
            ('PCs Genericos.xlsx', 4, {5: 'HOSTNAME'}, [
                (5, {3: 'prueba_local', 4: 'SECRETO_LOCAL', 5: 'pc_prueba'}),
            ]),
        ]
        for filename, header_row, headers, data_rows in specs:
            book = Workbook()
            sheet = book.active
            for column, text in headers.items():
                sheet.cell(header_row, column, text)
            for row_number, values in data_rows:
                for column, text in values.items():
                    sheet.cell(row_number, column, text)
            book.save(root / filename)

    def test_dry_run_is_read_only_and_report_omits_secrets(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            self.make_workbooks(root)
            report_path = root / 'report.json'
            output = io.StringIO()
            call_command('importar_datos_v11', input_dir=root, dry_run=True,
                         report=report_path, stdout=output)
            report_text = report_path.read_text(encoding='utf-8')
            report = json.loads(report_text)
            self.assertEqual(report['database_writes'], 0)
            self.assertEqual(report['planned']['usuarios'], 1)
            self.assertEqual(report['planned']['equipos'], 2)
            self.assertEqual(report['planned']['ips'], 1)
            self.assertEqual(len(report['equipos_sin_usuario']), 2)
            self.assertEqual(Usuario.objects.count(), 0)
            self.assertEqual(Equipamiento.objects.count(), 0)
            for secret in ('CLAVE_MUY_SECRETA', 'OTRA_CLAVE_SECRETA', 'SECRETO_LOCAL'):
                self.assertNotIn(secret, report_text)
                self.assertNotIn(secret, output.getvalue())

    def test_apply_creates_only_safe_relations_and_encrypts_secrets(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            self.make_workbooks(root)
            with self.assertRaises(CommandError):
                call_command('importar_datos_v11', input_dir=root, apply=True,
                             expected_sha256='wrong', stdout=io.StringIO())
            self.assertEqual(Usuario.objects.count(), 0)

            call_command('importar_datos_v11', input_dir=root, apply=True,
                         expected_sha256=file_digest(root), stdout=io.StringIO())

            user = Usuario.objects.get(usuario_red='prueba')
            self.assertEqual(user.departamento.nombre, 'TI')
            self.assertIsNone(user.subarea)
            self.assertEqual(user.gmail, 'aliasprueba@gmail.com')
            self.assertEqual(decrypt_val(user.password_gmail), 'CLAVE_MUY_SECRETA')
            equipment = Equipamiento.objects.get(numero_serie='SERIE123')
            self.assertIsNone(equipment.usuario)
            self.assertEqual(equipment.estado, 'STOCK')
            self.assertIsNone(Equipamiento.objects.get(numero_serie='SERIE_OK').usuario)
            ip = IP.objects.get(direccion_ip='192.168.10.10')
            self.assertEqual(ip.estado, 'LIBRE')
            self.assertIsNone(ip.usuario)
            self.assertIsNone(ip.observacion)
            self.assertEqual(Anexo.objects.get(numero_anexo='4000').usuario, user)
            self.assertIsNone(Anexo.objects.get(numero_anexo='4001').usuario)
            profile = PerfilGenerico.objects.get(usuario='perfil_prueba')
            self.assertEqual(profile.tipo, 'On Premise')
            self.assertEqual(decrypt_val(profile.password), 'OTRA_CLAVE_SECRETA')
            pc = PCGenerico.objects.get(hostname='pc_prueba')
            self.assertEqual(decrypt_val(pc.password), 'SECRETO_LOCAL')
            self.assertEqual(Departamento.objects.count(), 1)

            before = tuple(model.objects.count() for model in (
                Departamento, Usuario, Equipamiento, IP, Anexo,
                PerfilGenerico, PCGenerico,
            ))
            call_command('importar_datos_v11', input_dir=root, apply=True,
                         expected_sha256=file_digest(root), stdout=io.StringIO())
            after = tuple(model.objects.count() for model in (
                Departamento, Usuario, Equipamiento, IP, Anexo,
                PerfilGenerico, PCGenerico,
            ))
            self.assertEqual(before, after)
