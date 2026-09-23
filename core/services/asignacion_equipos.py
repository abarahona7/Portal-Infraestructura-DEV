"""Conciliación determinista de equipos existentes con usuarios del Excel."""

import hashlib
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import load_workbook

from core.models import Equipamiento, Usuario


EMPTY = {'', '-', 'N/A', 'N/I', 'FALTA', 'S/C', 'NAN', 'NONE'}
SPECS = (
    ('Notebook', 3, 4, 5, 6, 7),
    ('Celular', 8, 9, 10, 12, 15),
    ('Tablet', 17, 18, 19, 20, 23),
    ('BAM / Router', 25, 26, 27, 28, 30),
)


def clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = ' '.join(str(value).strip().split())
    return '' if text.upper() in EMPTY else text


def normalized(value):
    return clean(value).casefold()


def source_digest(path):
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _candidates(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        candidates = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            name = clean(row[1] if len(row) > 1 else None)
            for kind, brand_c, model_c, serial_c, af_c, date_c in SPECS:
                brand = clean(row[brand_c - 1] if len(row) >= brand_c else None)
                model = clean(row[model_c - 1] if len(row) >= model_c else None)
                serial = clean(row[serial_c - 1] if len(row) >= serial_c else None)
                af = clean(row[af_c - 1] if len(row) >= af_c else None)
                raw_date = row[date_c - 1] if len(row) >= date_c else None
                if any((brand, model, serial, af)):
                    candidates.append({
                        'row': row_number,
                        'kind': kind,
                        'name': name,
                        'brand': brand,
                        'model': model,
                        'serial': serial,
                        'af': af,
                        'assignment_date': _excel_date(raw_date),
                    })
    finally:
        book.close()
    return candidates


def _index(items, field):
    result = defaultdict(list)
    for item in items:
        value = normalized(getattr(item, field))
        if value:
            result[value].append(item.id)
    return result


def _pending(equipment, candidate, reason):
    return {
        'equipment_id': equipment.id,
        'row': candidate.get('row') if candidate else None,
        'kind': equipment.tipo,
        'brand': equipment.marca,
        'model': equipment.modelo,
        'serial': equipment.numero_serie or '',
        'af': equipment.af or '',
        'source_user': candidate.get('name', '') if candidate else '',
        'reason': reason,
    }


def build_plan(path):
    users = list(Usuario.objects.all())
    equipment = list(Equipamiento.objects.all())
    users_by_name = _index(users, 'nombre_completo')
    equipment_by_serial = _index(equipment, 'numero_serie')
    equipment_by_af = _index(equipment, 'af')
    user_by_id = {item.id: item for item in users}
    equipment_by_id = {item.id: item for item in equipment}

    groups = defaultdict(list)
    unmatched_candidates = []
    candidates = _candidates(path)

    for candidate in candidates:
        serial_ids = equipment_by_serial.get(normalized(candidate['serial']), [])
        af_ids = equipment_by_af.get(normalized(candidate['af']), [])
        matched_ids = set(serial_ids) | set(af_ids)
        conflict = bool(serial_ids and af_ids and set(serial_ids) != set(af_ids))
        if conflict or len(matched_ids) != 1:
            unmatched_candidates.append(candidate)
            continue
        candidate['user_ids'] = (
            users_by_name.get(normalized(candidate['name']), [])
            if candidate['name'] else []
        )
        groups[next(iter(matched_ids))].append(candidate)

    proposed = {}
    pending = []
    for item in equipment:
        matches = groups.get(item.id, [])
        if not matches:
            pending.append(_pending(item, None, 'EQUIPO_NO_IDENTIFICADO_EN_EL_EXCEL'))
            continue

        exact = [candidate for candidate in matches if len(candidate['user_ids']) == 1]
        destination_ids = {candidate['user_ids'][0] for candidate in exact}
        if len(destination_ids) > 1:
            pending.append(_pending(item, exact[0], 'USUARIOS_DESTINO_CONFLICTIVOS'))
            continue
        if not exact:
            representative = next((candidate for candidate in matches if candidate['name']), matches[0])
            if not representative['name']:
                reason = 'SIN_NOMBRE_DE_USUARIO'
            elif len(representative['user_ids']) > 1:
                reason = 'USUARIO_AMBIGUO'
            else:
                reason = 'USUARIO_SIN_COINCIDENCIA_EXACTA'
            pending.append(_pending(item, representative, reason))
            continue

        chosen = exact[0]
        user = user_by_id[chosen['user_ids'][0]]
        if item.usuario_id:
            reason = (
                'YA_ASIGNADO_AL_USUARIO_DEL_EXCEL'
                if item.usuario_id == user.id else
                'YA_ASIGNADO_A_OTRO_USUARIO'
            )
            pending.append(_pending(item, chosen, reason))
            continue
        if user.estado in {'BAJA', 'LICENCIA'}:
            pending.append(_pending(item, chosen, f'USUARIO_EN_ESTADO_{user.estado}'))
            continue
        proposed[item.id] = {
            'equipment': item,
            'user': user,
            'candidate': chosen,
        }

    notebook_hostnames = Counter(
        normalized(item['user'].hostname)
        for item in proposed.values()
        if item['equipment'].tipo in {'Notebook', 'Mac'} and item['user'].hostname
    )
    existing_hostnames = {
        normalized(value)
        for value in Equipamiento.objects.filter(
            usuario__isnull=False,
            tipo__in=['Notebook', 'Mac'],
        ).exclude(hostname__isnull=True).exclude(hostname='')
        .values_list('hostname', flat=True)
    }

    safe = []
    for equipment_id, item in proposed.items():
        hostname = normalized(item['user'].hostname)
        if (
            item['equipment'].tipo in {'Notebook', 'Mac'}
            and hostname
            and (notebook_hostnames[hostname] > 1 or hostname in existing_hostnames)
        ):
            pending.append(_pending(
                equipment_by_id[equipment_id],
                item['candidate'],
                'HOSTNAME_DE_NOTEBOOK_DUPLICADO',
            ))
            continue
        safe.append({
            'equipment_id': equipment_id,
            'user_id': item['user'].id,
            'assignment_date': item['candidate']['assignment_date'],
        })

    pending.sort(key=lambda item: (item['kind'], item['brand'], item['model'], item['serial']))
    return {
        'safe': safe,
        'pending': pending,
        'candidate_count': len(candidates),
        'unmatched_candidates': len(unmatched_candidates),
    }


def write_pending_report(path, plan, digest):
    lines = [
        'EQUIPOS EXISTENTES PENDIENTES DE ASIGNACIÓN',
        '=' * 120,
        f'Archivo origen: Usuarios Simi.xlsx',
        f'SHA-256: {digest}',
        f'Total pendiente: {len(plan["pending"])}',
        '',
        'N° | Fila Excel | Tipo | Marca / Modelo | N° Serie | Activo Fijo | Usuario informado | Motivo',
        '-' * 120,
    ]
    for index, item in enumerate(plan['pending'], 1):
        values = (
            index,
            item['row'] or 'N/A',
            item['kind'] or 'N/I',
            f'{item["brand"] or "N/I"} / {item["model"] or "N/I"}',
            item['serial'] or 'N/I',
            item['af'] or 'N/I',
            item['source_user'] or 'Sin usuario informado',
            item['reason'],
        )
        lines.append(' | '.join(str(value) for value in values))
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8-sig')
